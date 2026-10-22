#cell styles and  Named styles
#Cell styles are shared between objects and once they have been assigned they cannot be changed
import openpyxl
from openpyxl.styles import colors
from openpyxl.styles import Font, Color
from openpyxl import Workbook
wb = Workbook()
ws = wb.active

a1 = ws['A1']
d4 = ws['D4']
ft = Font(color="FF0000")
a1.font = ft
d4.font = ft

a1.font.italic = True # is not allowed 

# If you want to change the color of a Font, you need to reassign it::

a1.font = Font(color="FF0000", italic=True) # the change only affects A1



# copying styles between cells
from openpyxl.styles import Font
from copy import copy

ft1 = Font(name='Arial', size=14)
ft2 = copy(ft1)
ft2.name = "Tahoma"
ft1.name
'Arial'
ft2.name
'Tahoma'
ft2.size # copied from the
14.0

# colours 
# RGB hexadecimal color codes are used to specify colors in openpyxl. The format is RRGGBB, where RR is the red component, GG is the green component, and BB is the blue component. Each component is represented by two hexadecimal digits (00 to FF).
from openpyxl.styles import Font
font1 = Font(color="FF0000")

font2 = Font(color="00FF00")
font2.color.rgb #'0000FF00'


# There is also support for legacy indexed colours as well as themes and tints.
from openpyxl.styles.colors import Color
c = Color(indexed=32)
c = Color(theme=6, tint=0.5)


# Applying styles to a range of cells
from openpyxl.workbook import Workbook
from openpyxl.styles import Font, Fill
wb = Workbook()
ws = wb.active
c = ws['A1']
c.font = Font(size=12)

# applying a stles to columns and rows
col = ws.column_dimensions['A']
col.font = Font(bold=True)
row = ws.row_dimensions[1]
row.font = Font(underline="single")


# styling merged cells
from openpyxl.styles import Border, Side, PatternFill, Font, GradientFill, Alignment


wb = Workbook()
ws = wb.active
ws.merge_cells('B2:F4')

top_left_cell = ws['B2']
top_left_cell.value = "My Cell"

thin = Side(border_style="thin", color="000000")
double = Side(border_style="double", color="ff0000")

top_left_cell.border = Border(top=double, left=thin, right=thin, bottom=double)
top_left_cell.fill = PatternFill("solid", fgColor="DDDDDD")
top_left_cell.fill = fill = GradientFill(stop=("000000", "FFFFFF"))
top_left_cell.font  = Font(b=True, color="FF0000")
top_left_cell.alignment = Alignment(horizontal="center", vertical="center")

wb.save("styled.xlsx")


#using number formats
import datetime



ws = wb.active
# set date using a Python datetime
ws['A1'] = datetime.datetime(2010, 7, 21)

ws['A1'].number_format
'yyyy-mm-dd h:mm:ss'

ws["A2"] = 0.123456
ws["A2"].number_format = "0.00" # Display to 2dp



#creating a named style
from openpyxl.styles import NamedStyle, Font, Border, Side
highlight = NamedStyle(name="highlight")
highlight.font = Font(bold=True, size=20)
bd = Side(style='thick', color="000000")
highlight.border = Border(left=bd, top=bd, right=bd, bottom=bd)


wb.add_named_style(highlight)
ws['A1'].style = highlight
ws['D5'].style = 'highlight'