from openpyxl import Workbook

# create a new workbook
wb = Workbook()
wb.save("sample_work.xlsx")
# this allows you to access the default worksheet
ws = wb.active


# # create a new worksheet
# ws1 = wb.create_sheet(title="MySheet") # insert at the end (default)
# ws2 = wb.create_sheet(title="MyOtherSheet", index=0) # insert at first position
# ws3 = wb.create_sheet(title="MyThirdSheet", index=-1) # insert at penultimate position

# ws.title = "MyTitle" # rename the default sheet


ws["A1"] = "Name"
ws["B1"] = "Age"

ws["A2"] = "Alice"
ws["B2"] = 25

ws["A3"] = "Bob"
ws["B3"] = 30

wb.save("sample_work.xlsx")