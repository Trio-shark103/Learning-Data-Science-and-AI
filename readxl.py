import openpyxl

wb = openpyxl.load_workbook("sample_work.xlsx")
ws = wb.active

for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
    print(row)
    
# for row in ws.iter_rows(min_row=1, max_row=3):
#     for cell in row:
#         print(cell.value)
        
# for row in ws.iter_rows(min_row=1, max_row=3):
#     for cell in row:
#         print(cell.coordinate, cell.value)


for column in ws.iter_cols(min_col=1, max_col=2, min_row=1, max_row=3):
    for cell in column:
        print(cell.coordinate, cell.value)